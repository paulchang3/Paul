"""把 decompose.py 的結果寫成一份多分頁的 .xlsx。

分頁固定為(順序也固定,方便 Claude 或其他程式用分頁名稱直接讀取):
    00_Index / WBS_Tree / Action_Tree / Cornell_Notes / SMART_Notes / CODE_Notes

每個分頁第一列是表頭(粗體 + 凍結窗格 + 自動篩選),欄位順序固定,
00_Index 分頁用一份「欄位字典」說明每個分頁的用途與欄位定義,
讓沒看過這份檔案結構的人(或 LLM)也能直接看懂怎麼讀。
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

WBS_COLUMNS = ["level1", "level2", "level3", "level4", "path", "content", "source"]
ACTION_COLUMNS = ["task", "owner", "deadline", "priority", "status", "source_excerpt", "source"]
CORNELL_COLUMNS = ["section", "cues", "notes", "summary", "tags", "review_date", "source"]
SMART_COLUMNS = ["specific", "measurable", "achievable", "relevant", "deadline", "priority", "status", "source"]
CODE_COLUMNS = ["capture", "para", "tags", "distill", "express_type", "express", "source"]

SHEET_DEFS = [
    ("WBS_Tree", WBS_COLUMNS, "文件的知識樹狀結構(Work Breakdown Structure):"
     "level1~level4 是章節/條列的階層標題(愈右邊層級愈深,空白表示該列沒有更深層級),"
     "path 是完整階層路徑,content 是該節點下累積的原文內容。"),
    ("Action_Tree", ACTION_COLUMNS, "從文件中偵測到的行動/待辦項目:"
     "task 是行動內容,owner/deadline/priority/status 盡量從原文或表格擷取,"
     "抓不到的留空(需人工或 --llm 補齊)。"),
    ("Cornell_Notes", CORNELL_COLUMNS, "康乃爾筆記法:每一列是一個大章節,"
     "cues 是提示/問題、notes 是章節原文摘錄、summary 是精簡摘要、tags 是關鍵詞。"),
    ("SMART_Notes", SMART_COLUMNS, "SMART 目標筆記法:從行動項目或目標句子中,"
     "挑出含有可量化指標(數字/百分比/KPI)的項目;achievable/relevant 多半留白待人工填寫。"),
    ("CODE_Notes", CODE_COLUMNS, "CODE 筆記法(Capture/Organize/Distill/Express):"
     "capture 是原始摘錄、para 是 PARA 分類猜測、distill 是關鍵洞見、"
     "express/express_type 是建議的輸出形式與草稿。"),
]

HEADER_FILL_FONT = Font(bold=True)


def _write_sheet(wb: Workbook, name: str, columns: list[str], rows: list[dict]):
    ws = wb.create_sheet(name)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = HEADER_FILL_FONT
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append([row.get(col, "") for col in columns])
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    for i, col in enumerate(columns, start=1):
        width = max(12, min(60, len(col) + 4))
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_index(wb: Workbook, meta: dict):
    ws = wb.create_sheet("00_Index", 0)
    ws.append(["KM 知識拆解結果 — 分頁索引"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["產生時間", meta["generated_at"]])
    ws.append(["來源檔案", ", ".join(meta["sources"])])
    ws.append(["LLM 強化", "是(" + meta["llm_model"] + ")" if meta.get("llm_model") else "否(僅規則式)"])
    ws.append([])
    ws.append(["分頁名稱", "欄位", "說明"])
    for cell in ws[ws.max_row]:
        cell.font = HEADER_FILL_FONT
    for name, columns, desc in SHEET_DEFS:
        ws.append([name, ", ".join(columns), desc])
        ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 70


def write_workbook(result: dict, output_path: str | Path, sources: list[str],
                    llm_model: str | None = None) -> Path:
    output_path = Path(output_path)
    wb = Workbook()
    wb.remove(wb.active)  # 移除預設空白分頁

    meta = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "sources": sources,
        "llm_model": llm_model,
    }
    _write_index(wb, meta)
    _write_sheet(wb, "WBS_Tree", WBS_COLUMNS, result.get("wbs", []))
    _write_sheet(wb, "Action_Tree", ACTION_COLUMNS, result.get("actions", []))
    _write_sheet(wb, "Cornell_Notes", CORNELL_COLUMNS, result.get("cornell", []))
    _write_sheet(wb, "SMART_Notes", SMART_COLUMNS, result.get("smart", []))
    _write_sheet(wb, "CODE_Notes", CODE_COLUMNS, result.get("code", []))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
