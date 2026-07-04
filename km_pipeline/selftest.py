"""自我測試:產生一份合成的 HTML 範例,跑過整條流程,確認六個分頁都正確產生。

    python selftest.py

不需要 Ollama、不需要外部文件,幾秒內跑完。用來確認環境(openpyxl 等)
裝好了,以及改過程式碼後基本流程沒壞掉。
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

try:
    from . import build_excel, decompose, extract
except ImportError:
    import build_excel
    import decompose
    import extract

SAMPLE_HTML = """
<html><head><title>範例會議紀錄</title></head><body>
<h1>專案啟動會議</h1>
<p>本次會議討論 Q3 KM 系統上線計畫,與會人員包含專案經理與工程團隊。</p>
<h2>待辦事項</h2>
<ul>
  <li>負責人:小陳 完成 API 串接,截止 2026-08-01</li>
  <li>- [ ] 撰寫使用手冊</li>
</ul>
<h2>風險與資源</h2>
<p>參考文件:ISO 14971 規範。此為既有標準,不需修改。</p>
<table>
  <tr><th>項目</th><th>負責人</th><th>期限</th><th>狀態</th></tr>
  <tr><td>教育訓練</td><td>王小明</td><td>2026-08-15</td><td>進行中</td></tr>
</table>
</body></html>
"""


def main() -> int:
    checks: list[tuple[str, bool, str]] = []

    def check(name: str, ok: bool, detail: str = ""):
        checks.append((name, ok, detail))
        print(f"[{'OK' if ok else 'X '}] {name}" + (f" -- {detail}" if detail else ""))

    try:
        import openpyxl  # noqa: F401
        check("openpyxl 已安裝", True)
    except ImportError:
        check("openpyxl 已安裝", False, "pip install openpyxl")
        return 1

    with tempfile.TemporaryDirectory() as tmp:
        html_path = Path(tmp) / "sample.html"
        html_path.write_text(SAMPLE_HTML, encoding="utf-8")

        blocks = extract.extract(html_path)
        check("HTML 抽取出區塊", len(blocks) > 0, f"{len(blocks)} 個區塊")

        result = decompose.decompose(blocks, source=html_path.name)
        check("WBS 樹狀結構", len(result["wbs"]) > 0, f"{len(result['wbs'])} 節點")
        check("偵測到行動項目", len(result["actions"]) >= 2, f"{len(result['actions'])} 筆")
        check("康乃爾筆記", len(result["cornell"]) > 0, f"{len(result['cornell'])} 筆")
        check("CODE 筆記", len(result["code"]) > 0, f"{len(result['code'])} 筆")

        out_path = Path(tmp) / "out.xlsx"
        build_excel.write_workbook(result, out_path, sources=[html_path.name])
        check("Excel 已寫出", out_path.exists(), str(out_path))

        wb = openpyxl.load_workbook(out_path)
        expected_sheets = {"00_Index", "WBS_Tree", "Action_Tree", "Cornell_Notes",
                            "SMART_Notes", "CODE_Notes"}
        check("六個分頁都存在", expected_sheets.issubset(set(wb.sheetnames)),
              ", ".join(wb.sheetnames))

    passed = sum(1 for _, ok, _ in checks if ok)
    print(f"\n結果:{passed}/{len(checks)} 項通過")
    return 0 if passed == len(checks) else 1


if __name__ == "__main__":
    sys.exit(main())
